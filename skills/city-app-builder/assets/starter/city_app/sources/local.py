"""Read a chosen CSV or XLSX inside the app's private Data folder."""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

from .common import MAX_BYTES, SourceError, row_limit


def local_path(root: Path, filename: str) -> Path:
    if (
        not isinstance(filename, str)
        or not filename
        or Path(filename).is_absolute()
        or filename.startswith(("/", "\\"))
        or ":" in filename
    ):
        raise SourceError("Choose a file relative to this app's Data folder.")
    folder = (root / "Data").resolve()
    path = (folder / filename).resolve()
    if not path.is_relative_to(folder):
        raise SourceError("The input file must stay inside this app's Data folder.")
    if path.suffix.lower() not in (".csv", ".xlsx"):
        raise SourceError(
            "This connector accepts CSV or XLSX. Ask the user whether to convert this format or add a specific reader."
        )
    return path


def _headers(values) -> list[str]:
    if not values or any(not isinstance(v, str) or not v.strip() for v in values):
        raise SourceError("Use one header row with a nonempty name for every column.")
    headers = [v.strip() for v in values]
    if len(set(headers)) != len(headers):
        raise SourceError("Column headers must be unique.")
    return headers


def read_rows(root: Path, config: dict) -> list[dict]:
    path = local_path(root, config.get("path"))
    try:
        if path.stat().st_size > MAX_BYTES:
            raise SourceError(
                "The file exceeds the 10 MB limit. Supply a smaller extract."
            )
        if path.suffix.lower() == ".csv":
            with path.open("rb") as stream:
                raw = stream.read(MAX_BYTES + 1)
            if len(raw) > MAX_BYTES:
                raise SourceError("The file exceeds the 10 MB limit.")
            delimiter = config.get("delimiter", ",")
            if not isinstance(delimiter, str) or len(delimiter) != 1:
                raise SourceError("delimiter must be one character.")
            reader = csv.reader(
                io.StringIO(raw.decode("utf-8-sig")), delimiter=delimiter, strict=True
            )
            return _table_rows(reader, config)
        # Bound decompression before openpyxl parses workbook XML.
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50_000_000:
                raise SourceError(
                    "The expanded workbook exceeds 50 MB. Supply a smaller extract."
                )
        from openpyxl import load_workbook

        with path.open("rb") as stream:
            workbook = load_workbook(
                stream, read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = config.get("sheet")
                if sheet is None and len(workbook.sheetnames) != 1:
                    raise SourceError(
                        "This workbook has multiple sheets. Choose a sheet in the source configuration."
                    )
                if sheet is not None and sheet not in workbook.sheetnames:
                    raise SourceError("The selected workbook sheet does not exist.")
                worksheet = workbook[sheet or workbook.sheetnames[0]]
                worksheet.reset_dimensions()

                def values():
                    for cells in worksheet.iter_rows():
                        if any(cell.data_type == "f" for cell in cells):
                            raise SourceError(
                                "The worksheet contains formulas. Save a values-only copy so cached or uncalculated results cannot change totals."
                            )
                        if any(cell.data_type == "e" for cell in cells):
                            raise SourceError(
                                "The worksheet contains Excel errors. Correct them before loading."
                            )
                        yield [cell.value for cell in cells]

                return _table_rows(values(), config, pad_blanks=True)
            finally:
                workbook.close()
    except SourceError:
        raise
    except (
        OSError,
        UnicodeError,
        csv.Error,
        zipfile.BadZipFile,
        ValueError,
        KeyError,
    ) as error:
        raise SourceError(
            "Cannot read the local file. Check that it exists, is readable, and is valid UTF-8 CSV or XLSX."
        ) from error


def _table_rows(reader, config, *, pad_blanks=False):
    iterator = iter(reader)
    header_values = list(next(iterator, []))
    if pad_blanks:
        while header_values and (
            header_values[-1] is None or not str(header_values[-1]).strip()
        ):
            header_values.pop()
    headers = _headers(header_values)
    required = {name for name in config.get("columns", {}).values() if name is not None}
    if not required <= set(headers):
        raise SourceError(
            "The file header is missing a mapped column. Check the field mapping."
        )
    rows = []
    for values in iterator:
        values = list(values)
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if pad_blanks:
            while values and (values[-1] is None or not str(values[-1]).strip()):
                values.pop()
        if pad_blanks and len(values) < len(headers):
            values = [*values, *([None] * (len(headers) - len(values)))]
        if len(values) != len(headers):
            raise SourceError("A data row does not match the header width.")
        if len(rows) >= row_limit(config):
            raise SourceError(
                "The file exceeds max_rows. No partial totals were loaded."
            )
        rows.append(dict(zip(headers, values)))
    return rows
