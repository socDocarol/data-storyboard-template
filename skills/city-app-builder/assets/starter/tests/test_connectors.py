"""Connector contract tests using local files, a local HTTP server, and fake ODBC."""

import csv
import io
import json
import os
import tempfile
import threading
import unittest
from datetime import date, datetime
from decimal import Decimal
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from urllib.parse import parse_qs, urlsplit

from openpyxl import Workbook
from openpyxl.styles import PatternFill

import connect
from city_app.components import unavailable_panel
from city_app.data import export_csv, measure
from city_app.sources import arcgis, http, local, sql_server
from city_app.sources.common import SourceError, map_records
from city_app.sources.configured import ConfiguredSource, read_config

ROOT = Path(__file__).resolve().parents[1]


def config(name="local-csv"):
    return read_config(ROOT / "examples" / f"{name}.json")


ROW = {
    "record_id": "A",
    "date": "2025-01-02",
    "category": "Roads",
    "area": "North",
    "status": "Open",
    "value": "",
}


class LocalTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        (self.root / "Data").mkdir()
        self.config = config()

    def write_csv(self, rows, headers=None):
        with (self.root / "Data/example-services.csv").open(
            "w", encoding="utf-8-sig", newline=""
        ) as stream:
            writer = csv.DictWriter(stream, fieldnames=headers or list(ROW))
            writer.writeheader()
            writer.writerows(rows)

    def test_csv_mapping_missing_zero_credit_and_export(self):
        self.write_csv(
            [
                ROW,
                {**ROW, "record_id": "B", "value": "0", "area": ""},
                {**ROW, "record_id": "C", "value": "-2.5"},
            ]
        )
        dataset = ConfiguredSource(self.root, self.config).fetch()
        self.assertEqual([row.value for row in dataset.records], [None, 0, -2.5])
        self.assertEqual(dataset.records[1].area, "Not recorded")
        self.assertEqual(measure(dataset.records, dataset.source), (-1.25, 2))
        exported = list(
            csv.DictReader(io.StringIO(export_csv(dataset.records, dataset.source)))
        )
        self.assertEqual(exported[0]["data_kind"], "FICTIONAL SAMPLE")
        self.assertEqual(exported[2]["value"], "-2.5")

    def test_invalid_rows_fail_without_leaking_values(self):
        for change in (
            {"record_id": ""},
            {"date": "PRIVATE_BAD_DATE"},
            {"value": "nan"},
            {"value": True},
            {"category": []},
        ):
            with self.subTest(change=change), self.assertRaises(SourceError) as failure:
                map_records([{**ROW, **change}], self.config)
            self.assertNotIn("PRIVATE", str(failure.exception))
        with self.assertRaisesRegex(SourceError, "repeats"):
            map_records([ROW, ROW], self.config)
        with self.assertRaisesRegex(SourceError, "mapped column"):
            map_records([{"record_id": "A"}], self.config)

    def test_date_and_numeric_source_types(self):
        row = map_records(
            [{**ROW, "date": datetime(2025, 2, 3), "value": Decimal("4.5")}],
            self.config,
        )[0]
        self.assertEqual((row.date, row.value), (date(2025, 2, 3), 4.5))
        self.config["date_format"] = "%m/%d/%Y"
        self.assertEqual(
            map_records([{**ROW, "date": "02/03/2025"}], self.config)[0].date,
            date(2025, 2, 3),
        )
        self.config["date_format"] = "unix_ms"
        self.assertEqual(
            map_records([{**ROW, "date": 0}], self.config)[0].date, date(1970, 1, 1)
        )

    def test_explicit_row_count_is_distinct_from_missing_numeric_values(self):
        self.config["value_mode"] = "count"
        self.config["columns"]["value"] = None
        self.config["source"]["aggregation"] = "sum"
        self.assertEqual(map_records([ROW], self.config)[0].value, 1)
        self.config["value_mode"] = "column"
        self.assertIsNone(map_records([ROW], self.config)[0].value)

    def test_empty_missing_and_stale_are_real_connector_states(self):
        self.write_csv([])
        provider = ConfiguredSource(self.root, self.config)
        empty = provider.load("error")
        self.assertEqual(empty.state, "empty")
        self.assertNotIn("intentionally", str(unavailable_panel(empty)))
        self.write_csv([ROW])
        provider.checked_at = float("-inf")
        good = provider.load("loading")
        self.assertEqual(good.state, "ready")
        (self.root / "Data/example-services.csv").unlink()
        self.assertIs(provider.load(), good)
        provider.checked_at = float("-inf")
        stale = provider.load()
        self.assertEqual(stale.state, "stale")
        self.assertEqual(stale.records, good.records)
        self.assertEqual(stale.source.updated_at, good.source.updated_at)
        self.assertTrue(stale.message)
        failure = ConfiguredSource(self.root, self.config).load()
        self.assertEqual(failure.state, "error")
        self.assertNotIn("simulated", str(unavailable_panel(failure)))

    def test_file_and_row_budgets_and_private_paths(self):
        for path in (
            "../outside.csv",
            "C:/elsewhere.csv",
            "//server/share.csv",
            "book.xls",
            "file.json",
        ):
            with self.subTest(path=path), self.assertRaises(SourceError):
                local.local_path(self.root, path)
        self.write_csv([ROW, {**ROW, "record_id": "B"}])
        with self.assertRaisesRegex(SourceError, "max_rows"):
            local.read_rows(self.root, {**self.config, "max_rows": 1})
        with (
            patch("city_app.sources.local.MAX_BYTES", 20),
            self.assertRaisesRegex(SourceError, "MB"),
        ):
            local.read_rows(self.root, self.config)

    def test_csv_rejects_duplicate_headers_and_ragged_rows(self):
        for text in ("id,id\n1,2\n", "id,date\n1,2,3\n", "id,\n1,2\n"):
            (self.root / "Data/example-services.csv").write_text(text)
            with self.assertRaises(SourceError):
                local.read_rows(self.root, self.config)

    def test_xlsx_dates_sheet_selection_and_formula_rejection(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Requests"
        sheet.append(list(ROW))
        sheet.append(["A", datetime(2025, 1, 2), "Roads", "North", "Open", None])
        sheet["Z1"].fill = PatternFill("solid", fgColor="FFFF00")
        sheet["Z2"].fill = PatternFill("solid", fgColor="FFFF00")
        workbook.create_sheet("Notes")
        path = self.root / "Data/example-services.xlsx"
        workbook.save(path)
        cfg = config("local-xlsx")
        self.assertEqual(
            ConfiguredSource(self.root, cfg).fetch().records[0].date, date(2025, 1, 2)
        )
        with self.assertRaisesRegex(SourceError, "multiple sheets"):
            local.read_rows(self.root, {**cfg, "sheet": None})
        with self.assertRaisesRegex(SourceError, "does not exist"):
            local.read_rows(self.root, {**cfg, "sheet": "Missing"})
        sheet["F2"] = "=1+1"
        workbook.save(path)
        with self.assertRaisesRegex(SourceError, "formulas"):
            local.read_rows(self.root, cfg)
        workbook.close()

    def test_xlsx_does_not_discard_extra_data_without_a_header(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Requests"
        sheet.append(list(ROW))
        sheet.append([*ROW.values(), "unexpected extra value"])
        workbook.save(self.root / "Data/example-services.xlsx")
        workbook.close()
        with self.assertRaises(SourceError):
            local.read_rows(self.root, config("local-xlsx"))

    def test_inspection_is_bounded_and_does_not_echo_values(self):
        self.write_csv([{**ROW, "record_id": "PRIVATE_VALUE"}] * 120)
        profile = connect.inspect_file(self.root, "example-services.csv")
        self.assertEqual(profile["rows_inspected"], 100)
        self.assertNotIn("PRIVATE_VALUE", json.dumps(profile))

    def test_invalid_config_fails_before_a_source_is_opened(self):
        path = self.root / "config.json"
        for change in (
            {"source": []},
            {"kind": "unknown"},
            {"columns": {}},
            {"max_rows": 10001},
            {"refresh_seconds": 0},
        ):
            path.write_text(json.dumps({**self.config, **change}))
            with self.subTest(change=change), self.assertRaises(SourceError):
                read_config(path)


class ApiHandler(BaseHTTPRequestHandler):
    seen = []

    def do_POST(self):
        self.path += "?" + self.rfile.read(int(self.headers["Content-Length"])).decode()
        self.do_GET()

    def do_GET(self):
        parsed = urlsplit(self.path)
        query = parse_qs(parsed.query)
        self.seen.append((self.path, self.headers.get("Authorization")))
        if parsed.path == "/redirect":
            self.send_response(302)
            self.send_header("Location", "/must-not-be-called")
            self.end_headers()
            return
        if parsed.path == "/unauthorized":
            self.send_response(401)
            self.end_headers()
            return
        if parsed.path.endswith("/FeatureServer/0"):
            body = {
                "objectIdField": "id",
                "maxRecordCount": 1,
                "fields": [
                    {
                        "name": name,
                        "type": "esriFieldTypeOID"
                        if name == "id"
                        else "esriFieldTypeString",
                    }
                    for name in ("id", "opened", "group", "district", "status")
                ],
            }
        elif parsed.path.endswith("/query"):
            if "returnCountOnly" in query:
                body = {"count": 2}
            elif "returnIdsOnly" in query:
                body = {"objectIds": [1, 2]}
            else:
                identifier = int(query["objectIds"][0])
                body = {
                    "features": [
                        {
                            "attributes": {
                                "id": identifier,
                                "opened": 0,
                                "group": "Roads" if identifier == 1 else "Parks",
                                "district": 1,
                                "status": "Open",
                            }
                        }
                    ]
                }
        elif parsed.path == "/loop":
            body = {"records": [ROW], "next": "/loop"}
        elif parsed.path == "/cross-origin":
            body = {"records": [ROW], "next": "https://other.example/steal"}
        elif parsed.path == "/array":
            body = [ROW]
        else:
            index = int(query.get("page", query.get("offset", ["1"]))[0])
            rows = [{**ROW, "record_id": str(index)}] if index <= 2 else []
            body = {
                "records": rows,
                "next": f"/requests?page={index + 1}" if index < 2 else None,
            }
        raw = json.dumps(body).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def log_message(self, *args):
        pass


class HttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), ApiHandler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = f"http://127.0.0.1:{cls.server.server_port}"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()

    def setUp(self):
        ApiHandler.seen.clear()
        self.config = {**config("http-api"), "url": self.url + "/requests"}

    def test_json_array_and_all_pagination_modes(self):
        for mode in ("next", "page", "offset"):
            cfg = {
                **self.config,
                "pagination": {"mode": mode, "start": 1, "page_size": 1},
            }
            self.assertEqual(len(http.read_rows(cfg)), 2)
        cfg = {
            **self.config,
            "url": self.url + "/array",
            "records_path": "",
            "pagination": {"mode": "none"},
        }
        self.assertEqual(http.read_rows(cfg), [ROW])

    def test_loops_cross_origin_redirects_and_limits_fail_closed(self):
        for path, message in (
            ("/loop", "repeated"),
            ("/cross-origin", "different origin"),
            ("/redirect", "redirected"),
            ("/unauthorized", "401"),
        ):
            with self.subTest(path=path), self.assertRaisesRegex(SourceError, message):
                http.read_rows({**self.config, "url": self.url + path})
        self.assertFalse(
            any(path == "/must-not-be-called" for path, _ in ApiHandler.seen)
        )
        with self.assertRaisesRegex(SourceError, "max_rows"):
            http.read_rows({**self.config, "max_rows": 1})
        with (
            patch("city_app.sources.http.MAX_BYTES", 10),
            self.assertRaisesRegex(SourceError, "10 MB"),
        ):
            http.read_rows(self.config)

    def test_credentials_come_only_from_environment_and_are_not_in_errors(self):
        cfg = {**self.config, "auth": {"env": "TEST_STORYBOARD_TOKEN"}}
        with patch.dict(os.environ, {"TEST_STORYBOARD_TOKEN": "private-test-token"}):
            self.assertEqual(len(http.read_rows(cfg)), 2)
        self.assertEqual(
            {token for _, token in ApiHandler.seen}, {"Bearer private-test-token"}
        )
        with (
            patch.dict(os.environ, {"TEST_STORYBOARD_TOKEN": ""}),
            self.assertRaisesRegex(SourceError, "not set"),
        ):
            http.read_rows(cfg)
        for url in (
            "http://example.org/data",
            "https://user:password@example.org/data",
            "https://example.org/data?token=secret",
        ):
            with self.assertRaises(SourceError):
                http.validate_url(url)
        for key in (
            "client_secret",
            "subscription-key",
            "Ocp-Apim-Subscription-Key",
            "sig",
            "X-Amz-Signature",
            "auth",
        ):
            with self.subTest(key=key), self.assertRaises(SourceError):
                http.validate_url(f"https://example.org/data?{key}=private")
        http.validate_url("https://example.org/data?pageToken=public-cursor")

    def test_arcgis_live_protocol_fetches_every_id_in_batches(self):
        cfg = config("sacramento")
        cfg.update(
            url=self.url + "/FeatureServer/0",
            columns={
                "record_id": "id",
                "date": "opened",
                "category": "group",
                "area": "district",
                "status": "status",
                "value": None,
            },
        )
        dataset = ConfiguredSource(ROOT, cfg).fetch()
        self.assertEqual([row.record_id for row in dataset.records], ["1", "2"])
        self.assertTrue(all(row.date == date(1970, 1, 1) for row in dataset.records))
        self.assertEqual(sum("objectIds=" in path for path, _ in ApiHandler.seen), 2)
        self.assertFalse(dataset.source.is_sample)

    def test_arcgis_rejects_incomplete_download(self):
        cfg = config("sacramento")
        metadata = {
            "objectIdField": "OBJECTID",
            "maxRecordCount": 100,
            "fields": [
                {"name": name, "type": "esriFieldTypeString"}
                for name in cfg["columns"].values()
                if name
            ],
        }
        for page in ({"features": [], "exceededTransferLimit": True}, {"features": []}):
            with (
                patch.object(
                    http.JsonClient,
                    "get",
                    side_effect=[metadata, {"count": 1}, {"objectIds": [1]}, page],
                ),
                self.assertRaises(SourceError),
            ):
                arcgis.read_rows(cfg)

    def test_arcgis_hub_and_item_resolution(self):
        client = MagicMock()
        client.get.return_value = {
            "url": "https://services.arcgis.com/org/arcgis/rest/services/Data/FeatureServer"
        }
        layer = arcgis.resolve_layer(
            "https://data.cityofsacramento.org/datasets/" + "a" * 32 + "_3/about",
            client,
        )
        self.assertTrue(layer.endswith("FeatureServer/3"))
        client.get.side_effect = [
            {
                "url": "https://services.arcgis.com/org/arcgis/rest/services/Data/FeatureServer"
            },
            {"layers": [{"id": 0}, {"id": 1}]},
        ]
        with self.assertRaisesRegex(SourceError, "multiple"):
            arcgis.resolve_layer(
                "https://www.arcgis.com/home/item.html?id=" + "b" * 32, client
            )


class SqlTests(unittest.TestCase):
    def sql_fixture(self):
        cfg = {**config("sql-server"), "server": "work", "database": "reports"}
        cursor = MagicMock()
        cursor.description = [(name,) for name in cfg["columns"].values()]
        data = [("A", datetime(2025, 1, 2), "Roads", "North", "Open", Decimal("2.5"))]
        results = [
            [(1, "U", 0, 0)],
            [
                ("SERVER", "CONNECT SQL"),
                ("SERVER", "VIEW ANY DATABASE"),
                ("DATABASE", "CONNECT"),
                ("OBJECT", "SELECT"),
            ],
            [(name, 1, 0) for name in cfg["columns"].values()],
            data,
        ]
        cursor.fetchmany.side_effect = results
        connection = MagicMock()
        connection.cursor.return_value = cursor
        driver = SimpleNamespace(
            connect=MagicMock(return_value=connection),
            Error=RuntimeError,
            drivers=lambda: [sql_server.DRIVERS[1]],
        )
        return cfg, cursor, connection, driver, results

    def test_query_is_bounded_parameterized_and_identifiers_are_restricted(self):
        cfg = config("sql-server")
        query, parameters = sql_server.select_query(cfg)
        self.assertIn("SELECT TOP (1001)", query)
        self.assertIn("[OpenedDate] >= ?", query)
        self.assertNotIn("2025", query)
        self.assertEqual(parameters[0], date(2025, 1, 1))
        with self.assertRaises(SourceError):
            sql_server.select_query({**cfg, "table": "dbo.records;DROP TABLE x"})

    def test_optional_sql_driver_is_not_required_for_other_connectors(self):
        cfg = config("sql-server")
        with patch.dict("sys.modules", {"pyodbc": None}):
            result = sql_server.doctor(cfg)
            self.assertFalse(result["connection_attempted"])
            self.assertFalse(result["pyodbc_installed"])
            self.assertIsNone(result["selected_driver"])
            with self.assertRaisesRegex(SourceError, "app's Python environment"):
                sql_server.read_rows(cfg)

    def test_windows_auth_reuses_installed_driver_and_needs_no_secrets(self):
        cfg = {
            **config("sql-server"),
            "server": r"WORK\REPORTS",
            "database": "Reporting data",
        }
        for available, expected in (
            (list(sql_server.DRIVERS), sql_server.DRIVERS[0]),
            ([sql_server.DRIVERS[1]], sql_server.DRIVERS[1]),
        ):
            connection = sql_server.connection_string(cfg, available)
            self.assertIn(f"Driver={{{expected}}}", connection)
            self.assertIn(r"Server={WORK\REPORTS}", connection)
            self.assertIn("Database={Reporting data}", connection)
            self.assertIn("Trusted_Connection=yes;", connection)
            self.assertIn("Encrypt=yes;TrustServerCertificate=no;", connection)
            self.assertIn("ApplicationIntent=ReadOnly;", connection)
            self.assertNotIn("UID=", connection)
            self.assertNotIn("PWD=", connection)
        with self.assertRaisesRegex(SourceError, "does not install system drivers"):
            sql_server.connection_string(cfg, [])
        with self.assertRaisesRegex(SourceError, "Set server and database"):
            sql_server.connection_string(config("sql-server"), list(sql_server.DRIVERS))

    def test_sql_rejects_credentials_and_odbc_option_injection(self):
        cfg = {**config("sql-server"), "server": "work", "database": "reports"}
        for field in (
            "connection_env",
            "connection_string",
            "username",
            "password",
            "auth",
        ):
            with (
                self.subTest(field=field),
                self.assertRaisesRegex(SourceError, "Windows account"),
            ):
                sql_server.connection_string(
                    {**cfg, field: "not-allowed"}, list(sql_server.DRIVERS)
                )
        with self.assertRaises(sql_server.ReadOnlyViolation):
            sql_server.connection_string(
                {**cfg, "database": "reports};Trusted_Connection=no;PWD=bad"},
                list(sql_server.DRIVERS),
            )
        with self.assertRaises(SourceError):
            sql_server.connection_string(
                {**cfg, "server": "work\nserver"}, list(sql_server.DRIVERS)
            )

    def test_sql_doctor_only_checks_local_prerequisites(self):
        driver = SimpleNamespace(
            drivers=lambda: [sql_server.DRIVERS[1]], connect=MagicMock()
        )
        cfg = {**config("sql-server"), "server": "work", "database": "reports"}
        with patch.dict("sys.modules", {"pyodbc": driver}):
            result = sql_server.doctor(cfg)
        self.assertEqual(result["selected_driver"], sql_server.DRIVERS[1])
        self.assertTrue(result["server_configured"])
        self.assertTrue(result["database_configured"])
        self.assertFalse(result["connection_attempted"])
        driver.connect.assert_not_called()

    def test_sql_reads_and_closes_resources_without_real_database(self):
        cfg, cursor, connection, driver, results = self.sql_fixture()
        with patch.dict("sys.modules", {"pyodbc": driver}):
            rows = sql_server.read_rows(cfg)
        self.assertEqual(map_records(rows, cfg)[0].value, 2.5)
        self.assertTrue(driver.connect.call_args.kwargs["readonly"])
        self.assertFalse(driver.connect.call_args.kwargs["autocommit"])
        self.assertIn("Trusted_Connection=yes;", driver.connect.call_args.args[0])
        self.assertEqual(cursor.execute.call_args.args, sql_server.select_query(cfg))
        self.assertEqual(cursor.execute.call_count, 4)
        cursor.close.assert_called_once()
        connection.commit.assert_not_called()
        connection.rollback.assert_called_once()
        connection.close.assert_called_once()

    def test_sql_rejects_all_unsupported_settings_before_connecting(self):
        for key in (
            "query",
            "sql",
            "where",
            "init_sql",
            "stored_procedure",
            "connection_options",
            "readonly",
            "skip_permission_check",
            "UID",
            "Password",
            "allow_writes",
        ):
            cfg, cursor, connection, driver, _ = self.sql_fixture()
            cfg[key] = "DROP TABLE x"
            with (
                self.subTest(key=key),
                patch.dict("sys.modules", {"pyodbc": driver}),
                self.assertRaises(sql_server.ReadOnlyViolation),
            ):
                sql_server.read_rows(cfg)
            driver.connect.assert_not_called()

    def test_sql_identifiers_cannot_add_statements_functions_or_other_databases(self):
        for target in (
            "dbo.T;DELETE FROM T",
            "dbo.T--",
            "dbo.T/*x*/",
            "other.dbo.T",
            "server.db.dbo.T",
            "dbo.fn()",
            "dbo.[T]",
            "#Temp",
        ):
            cfg = {**config("sql-server"), "table": target}
            with (
                self.subTest(target=target),
                self.assertRaises(sql_server.ReadOnlyViolation),
            ):
                sql_server.select_query(cfg)
        for name in (
            "A];DROP TABLE x--",
            "COUNT(*)",
            "x INTO dbo.Copy",
            "NEXT VALUE FOR dbo.Sequence",
        ):
            cfg = config("sql-server")
            cfg["columns"]["value"] = name
            with (
                self.subTest(name=name),
                self.assertRaises(sql_server.ReadOnlyViolation),
            ):
                sql_server.select_query(cfg)

    def test_execution_boundary_refuses_writes_and_select_side_effects(self):
        cursor = MagicMock()
        valid, _ = sql_server.select_query(config("sql-server"))
        for statement in (
            "UPDATE dbo.T SET x=1",
            "DELETE FROM dbo.T",
            "INSERT INTO dbo.T VALUES (1)",
            "DROP TABLE dbo.T",
            "EXEC dbo.P",
            "SELECT * INTO dbo.Copy FROM dbo.T",
            "SELECT NEXT VALUE FOR dbo.S",
            "SELECT * FROM OPENROWSET('x')",
            valid + "; DELETE FROM dbo.T",
            valid + " -- comment",
            "SELECT TOP (99999) [Id] FROM [dbo].[T] ORDER BY [Id]",
        ):
            with (
                self.subTest(statement=statement),
                self.assertRaises(sql_server.ReadOnlyViolation),
            ):
                sql_server._execute_read(cursor, statement, ())
        cursor.execute.assert_not_called()

    def test_writable_or_unverifiable_permissions_block_data_query(self):
        for scope, permission in (
            ("SERVER", "CONTROL SERVER"),
            ("SERVER", "IMPERSONATE ANY LOGIN"),
            ("DATABASE", "INSERT"),
            ("DATABASE", "CREATE TABLE"),
            ("DATABASE", "EXECUTE"),
            ("SCHEMA", "ALTER"),
            ("OBJECT", "UPDATE"),
            ("OBJECT", "DELETE"),
            ("OBJECT", "CONTROL"),
            ("OBJECT", "TAKE OWNERSHIP"),
            ("OBJECT", None),
            ("OBJECT", "NEW UNKNOWN PERMISSION"),
        ):
            cfg, cursor, connection, driver, results = self.sql_fixture()
            results[1].append((scope, permission))
            with (
                self.subTest(scope=scope, permission=permission),
                patch.dict("sys.modules", {"pyodbc": driver}),
                self.assertRaises(sql_server.ReadOnlyViolation),
            ):
                sql_server.read_rows(cfg)
            self.assertEqual(cursor.execute.call_count, 2)
            connection.commit.assert_not_called()
            connection.rollback.assert_called_once()
            connection.close.assert_called_once()

    def test_admin_wrong_database_and_non_table_targets_are_rejected(self):
        for metadata in (
            [],
            [(0, "U", 0, 0)],
            [(1, "P", 0, 0)],
            [(1, "SN", 0, 0)],
            [(1, "U", 1, 0)],
            [(1, "U", 0, 1)],
            [(1, "U", None, 0)],
        ):
            cfg, cursor, connection, driver, results = self.sql_fixture()
            results[0][:] = metadata
            with (
                self.subTest(metadata=metadata),
                patch.dict("sys.modules", {"pyodbc": driver}),
                self.assertRaises(sql_server.ReadOnlyViolation),
            ):
                sql_server.read_rows(cfg)
            self.assertEqual(cursor.execute.call_count, 1)
            connection.rollback.assert_called_once()

    def test_column_update_grant_overrides_and_missing_select_are_rejected(self):
        for extra in (
            ("UnmappedColumn", 0, 1),
            ("UnmappedColumn", 0, None),
            ("RecordId", 0, 0),
            ("RecordId", None, 0),
        ):
            cfg, cursor, connection, driver, results = self.sql_fixture()
            results[2][:] = [row for row in results[2] if row[0] != extra[0]]
            results[2].append(extra)
            with (
                self.subTest(extra=extra),
                patch.dict("sys.modules", {"pyodbc": driver}),
                self.assertRaises(sql_server.ReadOnlyViolation),
            ):
                sql_server.read_rows(cfg)
            self.assertEqual(cursor.execute.call_count, 3)
            connection.rollback.assert_called_once()

    def test_permission_query_failure_stops_read_and_cleans_up(self):
        cfg, cursor, connection, driver, _ = self.sql_fixture()
        cursor.execute.side_effect = RuntimeError("private driver details")
        with (
            patch.dict("sys.modules", {"pyodbc": driver}),
            self.assertRaises(sql_server.ReadOnlyViolation) as failure,
        ):
            sql_server.read_rows(cfg)
        self.assertNotIn("private driver details", str(failure.exception))
        self.assertEqual(cursor.execute.call_count, 1)
        cursor.close.assert_called_once()
        connection.rollback.assert_called_once()
        connection.close.assert_called_once()

    def test_data_query_failure_and_row_cap_never_commit(self):
        for outcome in (RuntimeError("private driver details"), [(1,)] * 1001):
            cfg, cursor, connection, driver, results = self.sql_fixture()
            cursor.fetchmany.side_effect = [*results[:3], outcome]
            with (
                patch.dict("sys.modules", {"pyodbc": driver}),
                self.assertRaises(SourceError),
            ):
                sql_server.read_rows(cfg)
            connection.commit.assert_not_called()
            connection.rollback.assert_called_once()
            connection.close.assert_called_once()

    def test_cleanup_failure_refuses_result_and_still_attempts_rollback_and_close(self):
        for failure in ("cursor", "rollback", "connection"):
            cfg, cursor, connection, driver, _ = self.sql_fixture()
            operation = {
                "cursor": cursor.close,
                "rollback": connection.rollback,
                "connection": connection.close,
            }[failure]
            operation.side_effect = RuntimeError("private cleanup details")
            with (
                self.subTest(failure=failure),
                patch.dict("sys.modules", {"pyodbc": driver}),
                self.assertRaises(sql_server.ReadOnlyViolation) as error,
            ):
                sql_server.read_rows(cfg)
            self.assertNotIn("private cleanup details", str(error.exception))
            connection.commit.assert_not_called()
            connection.rollback.assert_called_once()
            connection.close.assert_called_once()

    def test_safety_refusal_clears_cached_sql_snapshot(self):
        cfg, cursor, connection, driver, _ = self.sql_fixture()
        provider = ConfiguredSource(ROOT, cfg)
        with patch.dict("sys.modules", {"pyodbc": driver}):
            self.assertEqual(provider.load().state, "ready")
        provider.checked_at = float("-inf")
        with patch.object(
            provider,
            "fetch",
            side_effect=sql_server.ReadOnlyViolation("Read-only access refused."),
        ):
            failure = provider.load()
        self.assertEqual(failure.state, "error")
        self.assertFalse(failure.records)


if __name__ == "__main__":
    unittest.main()
